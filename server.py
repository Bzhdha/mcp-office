"""MCP Office: convertit les réponses IA en DOCX/XLSX/PPTX à partir du dernier modèle d'entreprise (SharePoint)."""
from copy import deepcopy
import io,os,re,json,time,secrets,zipfile,hmac,hashlib,shutil,httpx,uvicorn
from pathlib import Path
from urllib.parse import quote
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook,Workbook
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn as pq
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from mcp.server.fastmcp import FastMCP,Context
from starlette.responses import Response,JSONResponse,HTMLResponse

E=os.environ.get
TENANT,CID,CSEC,SITE,FOLDER=E("SP_TENANT_ID"),E("SP_CLIENT_ID"),E("SP_CLIENT_SECRET"),E("SP_SITE_ID"),E("SP_FOLDER","Modeles")
LOCAL=E("TEMPLATE_DIR");API_KEY=E("MCP_API_KEY","");BASE=E("PUBLIC_BASE_URL","http://localhost:8000").rstrip("/")
OUT=Path(E("OUTPUT_DIR","/tmp/mcp-office"));OUT.mkdir(parents=True,exist_ok=True);TTL=int(E("FILE_TTL","3600"));MAXIN=int(E("MAX_INPUT","500000"))
LANG=E("DOC_LANG","fr-FR");CACHE=int(E("TEMPLATE_CACHE","300"))
CAT=Path(E("CATALOG_DIR","/data/catalog"));MODELS=Path(E("MODELS_DIR",str(Path(__file__).with_name("catalog"))));BUNDLE=MODELS/"bundle"
EXT={"docx":(".dotx",".docx"),"xlsx":(".xltx",".xlsx"),"pptx":(".potx",".pptx")}
MIME={"docx":"application/vnd.openxmlformats-officedocument.wordprocessingml.document","xlsx":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet","pptx":"application/vnd.openxmlformats-officedocument.presentationml.presentation"}
G="https://graph.microsoft.com/v1.0";_tok=[None,0];_tpl={}

# ---------- Modèles (SharePoint via Graph, ou dossier local synchronisé) ----------
def _gtoken():
 if _tok[1]>time.time()+60:return _tok[0]
 r=httpx.post(f"https://login.microsoftonline.com/{TENANT}/oauth2/v2.0/token",data={"grant_type":"client_credentials","client_id":CID,"client_secret":CSEC,"scope":"https://graph.microsoft.com/.default"},timeout=20);r.raise_for_status();j=r.json()
 _tok[:]=[j["access_token"],time.time()+j["expires_in"]];return _tok[0]
def _list():
 if not LOCAL and not all((TENANT,CID,CSEC,SITE)):raise ValueError("Aucune source de modèles configurée (SP_* ou TEMPLATE_DIR) : le modèle PowerPoint épinglé reste utilisé")
 if LOCAL:return[{"name":p.name,"modified":p.stat().st_mtime,"id":str(p),"etag":str(p.stat().st_mtime)}for p in Path(LOCAL).iterdir()if p.is_file()]
 h={"Authorization":"Bearer "+_gtoken()};u=f"{G}/sites/{SITE}/drive/root:/{FOLDER.strip('/')}:/children?$select=id,name,lastModifiedDateTime,eTag,file&$top=999";items=[]
 while u:r=httpx.get(u,headers=h,timeout=20);r.raise_for_status();j=r.json();items+=j["value"];u=j.get("@odata.nextLink")
 return[{"name":i["name"],"modified":i["lastModifiedDateTime"],"id":i["id"],"etag":i["eTag"]}for i in items if"file"in i]
def latest(kind,prefix=""):
 """Dernier modèle modifié pour le type (préfixe de nom optionnel)."""
 c=[f for f in _list()if f["name"].lower().endswith(EXT[kind])and not f["name"].startswith("~$")and f["name"].lower().startswith(prefix.lower())]
 if not c:raise ValueError(f"Aucun modèle {kind} trouvé")
 f=max(c,key=lambda x:x["modified"]);k=(kind,prefix);hit=_tpl.get(k)
 if hit and hit[0]==f["etag"]and hit[2]>time.time():return f["name"],hit[1]
 if LOCAL:b=Path(f["id"]).read_bytes()
 else:r=httpx.get(f"{G}/sites/{SITE}/drive/items/{f['id']}/content",headers={"Authorization":"Bearer "+_gtoken()},follow_redirects=True,timeout=60);r.raise_for_status();b=r.content
 _tpl[k]=(f["etag"],b,time.time()+CACHE);return f["name"],b
def _untemplate(b):
 """.dotx/.potx -> document ouvrable (content-type template -> document)."""
 zi,o=zipfile.ZipFile(io.BytesIO(b)),io.BytesIO()
 with zipfile.ZipFile(o,"w",zipfile.ZIP_DEFLATED)as z:
  for it in zi.infolist():
   d=zi.read(it)
   if it.filename=="[Content_Types].xml":d=d.replace(b"wordprocessingml.template.main",b"wordprocessingml.document.main").replace(b"presentationml.template.main",b"presentationml.presentation.main")
   z.writestr(it,d)
 o.seek(0);return o

# ---------- Markdown minimal ----------
def blocks(md):
 L=md.replace("\r","").split("\n");i=0;out=[]
 while i<len(L):
  l=L[i];s=l.strip()
  if not s:i+=1;continue
  if s.startswith("```"):
   i+=1;c=[]
   while i<len(L)and not L[i].strip().startswith("```"):c.append(L[i]);i+=1
   out.append(("code","\n".join(c)));i+=1;continue
  if m:=re.match(r"(#{1,6})\s+(.*)",s):out.append(("h",len(m[1]),m[2]));i+=1;continue
  if s.startswith("|"):
   rows=[]
   while i<len(L)and L[i].strip().startswith("|"):
    r=[c.strip()for c in L[i].strip().strip("|").split("|")]
    if not all(re.fullmatch(r":?-{2,}:?",c)for c in r if c):rows.append(r)
    i+=1
   out.append(("table",rows));continue
  if m:=re.match(r"(\s*)([-*+]|\d+[.)])\s+(.*)",l):out.append(("num"if m[2][0].isdigit()else"ul",min(len(m[1].replace("\t","  "))//2,2),m[3]));i+=1;continue
  p=[s];i+=1
  while i<len(L)and L[i].strip()and not re.match(r"\s*(#|\||```|[-*+]\s|\d+[.)]\s)",L[i]):p.append(L[i].strip());i+=1
  out.append(("p"," ".join(p)))
 return out
def runs(t):return[(x.strip("*`"),x.startswith("**"),x.startswith("*")and not x.startswith("**"),x.startswith("`"))for x in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)",t)if x]
def plain(t):return"".join(r[0]for r in runs(t))

# ---------- Génération ----------
def _save(kind,name,obj):
 for f in OUT.iterdir():
  if f.stat().st_mtime<time.time()-TTL:shutil.rmtree(f,ignore_errors=True)if f.is_dir()else f.unlink(missing_ok=True)
 tok=secrets.token_urlsafe(24);safe=re.sub(r"[^\w\-. ]","_",name)[:80].strip()or"document"
 d=OUT/tok;d.mkdir();obj.save(d/f"{safe}.{kind}");return f"{BASE}/files/{tok}/{quote(safe)}.{kind}"

def make_docx(title,markdown,template_prefix=""):
 tn,b=latest("docx",template_prefix);doc=Document(_untemplate(b));body=doc.element.body
 for e in list(body):
  if e.tag!=qn("w:sectPr"):body.remove(e)
 names={s.name for s in doc.styles};st=lambda*n:next((x for x in n if x in names),None)
 doc.core_properties.title=title;doc.core_properties.language=LANG
 def para(t,style=None,pre=""):
  p=doc.add_paragraph(style=style);pre and p.add_run(pre)
  for x,b_,i_,c in runs(t):r=p.add_run(x);r.bold=b_ or None;r.italic=i_ or None;c and setattr(r.font,"name","Consolas")
  return p
 if title:doc.add_heading(title,0)if st("Title")else para(title,st("Heading 1"))
 for bl in blocks(markdown):
  k=bl[0]
  if k=="h":doc.add_heading(plain(bl[2]),min(bl[1],9))
  elif k in("ul","num"):
   base="List Bullet"if k=="ul"else"List Number";s=st(base+(f" {bl[1]+1}"if bl[1]else""),base,"List Paragraph")
   para(bl[2],s,""if s and s.startswith(base)else("• "if k=="ul"else"- "))
  elif k=="code":para(bl[1],st("Code","HTML Preformatted"))
  elif k=="table":
   rows=bl[1];nc=max(map(len,rows));t=doc.add_table(rows=0,cols=nc);t.style=st("Table Grid","Light Grid Accent 1")
   for ri,r in enumerate(rows):
    cells=t.add_row().cells
    for ci in range(nc):
     c=cells[ci];c.text=plain(r[ci])if ci<len(r)else""
     if ri==0:
      for rn in c.paragraphs[0].runs:rn.bold=True
   h=t.rows[0]._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");h.append(e)  # en-tête répété (accessibilité)
  else:para(bl[1])
 return tn,doc

def _num(v):
 if isinstance(v,str):
  s=v.strip().replace(" ","").replace(" ","")
  if re.fullmatch(r"-?\d+([.,]\d+)?",s):return float(s.replace(",","."))if re.search(r"[.,]",s)else int(s)
 return v
def make_xlsx(title,sheets,template_prefix=""):
 tn,b=latest("xlsx",template_prefix);wb=load_workbook(io.BytesIO(b));wb.template=False;base=wb.worksheets[0]
 wb.properties.title=title;wb.properties.language=LANG;used=set()
 for n,sh in enumerate(sheets):
  ws=base if n==0 else wb.copy_worksheet(base);nm=re.sub(r"[\[\]:*?/\\]","",str(sh.get("name")or f"Feuille{n+1}"))[:31]or f"Feuille{n+1}"
  while nm in used:nm=nm[:28]+f"_{n}"
  used.add(nm);ws.title=nm;rows=sh.get("rows")or[];start=ws.max_row+2 if ws.max_row>1 or ws["A1"].value is not None else 1
  for ri,r in enumerate(rows):
   for ci,v in enumerate(r):ws.cell(start+ri,ci+1,_num(v))
  if len(rows)>1:
   nc=max(map(len,rows));ref=f"A{start}:{get_column_letter(nc)}{start+len(rows)-1}"
   hdr=[str(rows[0][i])if i<len(rows[0])and rows[0][i]not in(None,"")else f"Col{i+1}"for i in range(nc)]
   for i,h in enumerate(hdr):ws.cell(start,i+1,h)
   t=Table(displayName=re.sub(r"\W","_",f"T_{nm}")[:250],ref=ref);t.tableStyleInfo=TableStyleInfo(name=E("XLSX_TABLE_STYLE","TableStyleMedium2"),showRowStripes=True);ws.add_table(t)
   ws.freeze_panes=ws.cell(start+1,1)
   for i in range(nc):ws.column_dimensions[get_column_letter(i+1)].width=min(60,max(10,*(len(str(r[i]))+2 for r in rows if i<len(r))))
 for ws in wb.worksheets[len(sheets):]:wb.remove(ws)
 return tn,wb

# PowerPoint : le modèle contient des diapos types préparées -> on les duplique et on remplit leurs zones.
RNS="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
def _stitle(s):
 t=s.shapes.title
 return(t.text_frame.text.strip()if t is not None and t.has_text_frame else"")or s.slide_layout.name
def _pick(p,t,has_table=False):
 S=list(p.slides)
 if isinstance(t,int)or str(t).isdigit():return S[int(t)-1]
 if t:
  for s in S:
   if str(t).lower()in _stitle(s).lower():return s
  raise ValueError(f"Diapo type « {t} » introuvable (voir list_slide_types)")
 for s in S[1:]or S:  # auto : 1re diapo (hors couverture) avec tableau ou corps de texte
  if has_table and any(sh.has_table for sh in s.shapes):return s
  if not has_table and any(sh.is_placeholder and int(sh.placeholder_format.type)in(2,7)for sh in s.shapes):return s
 return S[-1]
def _dup(p,src):
 s=p.slides.add_slide(src.slide_layout);tree=s.shapes._spTree
 for sh in list(s.shapes):tree.remove(sh._element)
 m={}
 for rel in list(src.part.rels._rels.values()):
  if rel.reltype.endswith(("/slideLayout","/notesSlide")):continue
  m[rel.rId]=s.part.rels.get_or_add_ext_rel(rel.reltype,rel.target_ref)if rel.is_external else s.part.relate_to(rel.target_part,rel.reltype)
 return _copy_shapes(s,src,m)
def _import(p,src):
 """Copie une diapo d'une AUTRE présentation (diapo utilisateur) : même disposition du modèle, images et liens externes recopiés."""
 s=p.slides.add_slide(_layout(p,src.slide_layout.name));tree=s.shapes._spTree
 for sh in list(s.shapes):tree.remove(sh._element)
 m={}
 for rel in list(src.part.rels._rels.values()):
  if rel.reltype.endswith(("/slideLayout","/notesSlide")):continue
  if rel.is_external:m[rel.rId]=s.part.relate_to(rel.target_ref,rel.reltype,is_external=True)
  elif rel.reltype==RT.IMAGE:
   try:m[rel.rId]=s.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))[1]
   except Exception:pass  # format d'image non pris en charge (signalé à l'analyse)
 return _copy_shapes(s,src,m)
def _copy_shapes(s,src,m):
 tree=s.shapes._spTree
 bg=src._element.cSld.bg
 if bg is not None:s._element.cSld.insert(0,deepcopy(bg))
 for el in list(src.shapes._spTree)[2:]:
  e=deepcopy(el)
  for x in e.iter():
   for a in("embed","id","link"):
    v=x.get(f"{{{RNS}}}{a}")
    if v in m:x.set(f"{{{RNS}}}{a}",m[v])
  tree.append(e)
 return s
def _fonts(p):
 """Polices du thème (titres, texte), ex. ("N27 Medium","N27 Light")."""
 b=p.slide_masters[0].part.part_related_by(RT.THEME).blob.decode("utf8","ignore")
 f=[re.search(rf"<a:{k}Font>\s*<a:latin typeface=\"([^\"]*)\"",b)for k in("major","minor")]
 return tuple(m[1]if m else""for m in f)
def _bu(p,t):pr=p.find(pq("a:pPr"));return pr is not None and any(pr.find(pq(x))is not None for x in t)
def _bul(p,explicit=True):
 """Paragraphe à puce : puce explicite, ou (si l'exemple n'en a aucune d'explicite) puce héritée non annulée par buNone."""
 return _bu(p,("a:buChar","a:buAutoNum","a:buBlip"))if explicit else not _bu(p,("a:buNone",))
def _fill(tf,lines,fonts=("","")):
 """Remplace le texte : la ligne n reprend la mise en forme du paragraphe n de l'exemple (le dernier au-delà).
 Si des lignes commencent par « - », structure respectée : lignes « - » -> paragraphes à puce de l'exemple, autres -> paragraphes sans puce, "" -> paragraphe vide.
 Gras = police titres du thème si distincte (charte : mots importants en N27 Medium) ; polices hors thème retirées ; \\v = saut de ligne."""
 lines=lines or[""];ps=tf._txBody.findall(pq("a:p"));full=[x for x in ps if"".join(x.itertext()).strip()]or ps[:1]
 fam=(fonts[1]or fonts[0]).split(" ")[0];cls=any(re.match(r"\s*[-*•]\s",l)for l in lines);seen={}
 for x in ps:tf._txBody.remove(x)
 for n,ln in enumerate(lines):
  if not cls:p0=full[min(n,len(full)-1)]
  else:
   ex=any(_bul(x)for x in full)
   k="e"if not ln.strip()else"b"if re.match(r"\s*[-*•]\s",ln)else"t";c=[x for x in(ps if k=="e"else full)if(k=="e")==(not"".join(x.itertext()).strip())and(k=="e"or _bul(x,ex)==(k=="b"))]or full
   p0=c[min(seen.get(k,0),len(c)-1)];seen[k]=seen.get(k,0)+1
  p=deepcopy(p0);rs=[r.find(pq("a:rPr"))for r in p0.findall(pq("a:r"))];rs=[r for r in rs if r is not None]
  rpr=next((r for r in rs if not any(f.get("typeface")==fonts[0]for f in r.findall(pq("a:latin")))),rs[0]if rs else None)  # éviter de reprendre la police d'accent
  for c in list(p):
   if c.tag!=pq("a:pPr"):p.remove(c)
  tf._txBody.append(p);pa=tf.paragraphs[-1];lvl=(len(ln)-len(ln.lstrip(" ")))//2
  if lvl:pa.level=min(lvl,8)
  for k,seg in enumerate(re.sub(r"^[-*•]\s+","",ln.strip()).split("\v")):
   k and pa.add_line_break()
   for x,b_,i_,_ in runs(seg):
    r=pa.add_run()
    if rpr is not None:
     q=deepcopy(rpr);r._r.insert(0,q)
     for f in q.findall(pq("a:latin")):
      if fam and not f.get("typeface","+").startswith(("+",fam)):q.remove(f)
    r.text=x;i_ and setattr(r.font,"italic",True)
    if b_:
     if fonts[0]and fonts[0]!=fonts[1]:r.font.name=fonts[0]
     else:r.font.bold=True
def _table(t,rows,fonts=("","")):
 nc=len(t.columns)
 while len(t._tbl.tr_lst)<len(rows):t._tbl.append(deepcopy(t._tbl.tr_lst[-1]))
 for x in t._tbl.tr_lst[len(rows):]:t._tbl.remove(x)
 for ri,r in enumerate(rows):
  for ci in range(nc):_fill(t.cell(ri,ci).text_frame,[str(r[ci])if ci<len(r)else""],fonts)
def _layout(p,name):
 L=[l for m in p.slide_masters for l in m.slide_layouts]
 for ok in(lambda l:l.name.lower()==str(name).lower(),lambda l:str(name).lower()in l.name.lower()):
  for l in L:
   if ok(l):return l
 raise ValueError(f"Disposition « {name} » introuvable (voir list_slide_types(layouts=True))")
def _zone(s,key):
 """Zone par nom (toutes les formes de ce nom), « Nom#n » (n-ième forme de ce nom) ou « @idx » (espace réservé d'index idx)."""
 if m:=re.fullmatch(r"@(\d+)",key):return[sh for sh in s.shapes if sh.is_placeholder and sh.placeholder_format.idx==int(m[1])]
 n,_,i=key.rpartition("#")
 if n and i.isdigit():return[sh for sh in s.shapes if sh.name==n][int(i)-1:int(i)]
 return[sh for sh in s.shapes if sh.name==key]
# ---------- Catalogue PowerPoint : modèle épinglé + analyse enregistrée + diapos autorisées ----------
# État (CATALOG_DIR, volume) : copie du modèle épinglé et analyse (polices, inventaire des zones). Refaite seulement via refresh_template.
# Configuration (MODELS_DIR, versionnée) : modèles de diapos exposés au chat (champs nommés, limites, textes fixes, règles).
def _cfile(d,prefix,ext=".json"):return Path(d)/("pptx"+("-"+re.sub(r"\W","_",prefix)if prefix else"")+ext)
_cats={}
def _inventory(p):
 inv={"slides":[],"layouts":[]}
 for i,s in enumerate(p.slides,1):
  seen={};names=[sh.name for sh in s.shapes];z=[]
  for sh in s.shapes:
   seen[sh.name]=seen.get(sh.name,0)+1;k=sh.name if names.count(sh.name)==1 else f"{sh.name}#{seen[sh.name]}"
   kind="table"if sh.has_table else"title"if sh.is_placeholder and int(sh.placeholder_format.type)in(1,3)else"text"if sh.has_text_frame else None
   if kind:z.append({"zone":k,"kind":kind,"exemple":sh.text_frame.text[:80]if sh.has_text_frame else f"{len(sh.table.rows)}x{len(sh.table.columns)}"})
  inv["slides"].append({"type":i,"label":_stitle(s),"layout":s.slide_layout.name,"zones":z})
 sigs={}
 for mi,m in enumerate(p.slide_masters):
  for l in m.slide_layouts:
   if any(x["layout"]==l.name for x in inv["layouts"]):continue
   ph=[h for h in l.placeholders if int(h.placeholder_format.type)not in(13,15,16)]
   inv["layouts"].append({"layout":l.name,"master":mi+1,"zones":[{"zone":f"@{h.placeholder_format.idx}","kind":"title"if int(h.placeholder_format.type)in(1,3)else"image"if int(h.placeholder_format.type)==18 else"text","exemple":h.text_frame.text[:40]}for h in ph]})
   if ph:sigs.setdefault(tuple(sorted((h.placeholder_format.idx,int(h.placeholder_format.type),round(h.left/36e4),round(h.top/36e4))for h in ph)),[]).append(l.name)
 inv["families"]={}  # variantes = dispositions de même structure (mêmes zones, mêmes positions), visuel différent
 for v in sigs.values():
  if len(v)<2:continue
  n=os.path.commonprefix(v).rstrip(" -_");n=n if len(n)>3 else v[0]
  while n in inv["families"]:n+="+"
  inv["families"][n]=v
 return inv
def _auto_models(inv):
 """Sans configuration : une entrée par diapo type, champs = zones brutes."""
 return{f"diapo_{s['type']}":{"usage":s["label"],"source":{"type":s["type"]},"fields":{re.sub(r"\W+","_",z["zone"]).strip("_").lower():{"zone":z["zone"],"kind":"table"if z["kind"]=="table"else"text","exemple":z["exemple"]}for z in s["zones"]}}for s in inv["slides"]}
def _variants(m,inv):
 """Variantes d'un modèle de diapo : [{"type":n} | {"layout":nom}, + label/ambiance de la config]."""
 s=m["source"];info=m.get("variants")or{}
 if"family"in s:v=[{"layout":x}for x in inv["families"].get(s["family"],[])]
 else:
  k="type"if"type"in s else"layout";v=[{k:x}for x in(s[k]if isinstance(s[k],list)else[s[k]])]
 return[{**x,**info.get(str(next(iter(x.values()))),{})}for x in v]
def _check(c,p):
 w=[];S=list(p.slides);lay={l["layout"]for l in c["inventory"]["layouts"]}
 for n,m in c["models"].items():
  keys=[f["zone"]for f in(m.get("fields")or{}).values()]+list((m.get("fixed")or{}).keys())
  if not m["_v"]:w.append(f"{n}: aucune variante (famille ou source introuvable)")
  for v in m["_v"]:
   try:
    if"type"in v:s=S[int(v["type"])-1];w+=[f"{n}: zone {k} absente de la diapo {v['type']}"for k in keys if not _zone(s,k)]
    else:
     L=_layout(p,v["layout"]);ok={f"@{h.placeholder_format.idx}"for h in L.placeholders}
     w+=[f"{n}: zone {k} absente de la disposition {v['layout']}"for k in keys if k not in ok]
   except Exception as e:w.append(f"{n}: source {v} invalide ({e})")
 used={v.get("layout")for m in c["models"].values()for v in m["_v"]}|{S[int(v["type"])-1].slide_layout.name for m in c["models"].values()for v in m["_v"]if"type"in v and int(v["type"])<=len(S)}
 libres=sorted(lay-used-set(c.get("ignored")or{}))
 if libres:w.append("dispositions non exposées au chat (à ajouter dans catalog/pptx.json si utiles) : "+", ".join(libres))
 return w
def catalog(prefix="",refresh=False):
 """Catalogue en cache ; l'analyse du modèle n'est refaite que si refresh=True (ou au premier usage)."""
 if prefix in _cats and not refresh:return _cats[prefix]
 sf=_cfile(CAT,prefix);st=json.loads(sf.read_text("utf8"))if sf.is_file()else{}
 cf=_cfile(MODELS,prefix);cfg=json.loads(cf.read_text("utf8"))if cf.is_file()else{}
 bf=_cfile(BUNDLE,prefix)
 if not refresh and not st.get("template")and bf.is_file()and BUNDLE!=CAT:  # 1er démarrage : modèle et analyse livrés dans l'image (python server.py bundle)
  bs=json.loads(bf.read_text("utf8"))
  if(BUNDLE/bs.get("template","")).is_file():CAT.mkdir(parents=True,exist_ok=True);(CAT/bs["template"]).write_bytes((BUNDLE/bs["template"]).read_bytes());sf.write_text(bf.read_text("utf8"),"utf8");st=bs
 if refresh or not st.get("template")or not(CAT/st["template"]).is_file():
  name,b=latest("pptx",prefix);CAT.mkdir(parents=True,exist_ok=True);(CAT/name).write_bytes(b);p=Presentation(_untemplate(b))
  st={"template":name,"analyzed":time.strftime("%Y-%m-%d %H:%M"),"fonts":dict(zip(("accent","texte"),_fonts(p))),"inventory":_inventory(p)}
  sf.write_text(json.dumps(st,ensure_ascii=False,indent=1),"utf8")
 b=(CAT/st["template"]).read_bytes()
 c={**st,"fonts":cfg.get("fonts")or st["fonts"],"rules":cfg.get("rules",[]),"groups":cfg.get("groups",[]),"ambiances":cfg.get("ambiances",{}),"default_ambiance":cfg.get("default_ambiance",""),"ignored":cfg.get("ignored_layouts",{}),"conformite":cfg.get("conformite",{}),"models":cfg.get("models")or _auto_models(st["inventory"]),"bytes":b}
 for m in c["models"].values():m["_v"]=_variants(m,st["inventory"])
 c["warnings"]=_check(c,Presentation(_untemplate(b)));_cats[prefix]=c;return c
def _pickv(c,n,m,d,used,amb):
 """Variante : champ pilote (variant_by), « variante » explicite, sinon alternance dans l'ambiance couleur du document."""
 V=m["_v"];vals=d.get("fields")or{}
 if m.get("variant_by"):return[V[min(max(int(vals.get(m["variant_by"],1))-1,0),len(V)-1)]]
 if m.get("sequence"):return V
 if d.get("variante"):return[V[min(max(int(d["variante"])-1,0),len(V)-1)]]
 pool=[x for x in V if x.get("ambiance")==amb]or V;k=used.get(n,0);used[n]=k+1
 return[pool[k%len(pool)]]
def _expand(c,d,used,warn,amb=""):
 """Diapo {"model","fields"} -> [{"type"|"layout","zones"}] + contrôle des champs (longueurs, nombre de lignes)."""
 n=d.get("model");m=c["models"].get(n)
 if not m:raise ValueError(f"Modèle de diapo inconnu « {n} » (voir get_presentation_catalog)")
 vs=_pickv(c,n,m,d,used,amb)
 SK=("type","layout","user","slide")
 if m.get("asis"):return[{**{k:v[k]for k in SK if k in v},"asis":True,"notes":d.get("notes")}for v in vs]
 src={k:vs[0][k]for k in SK if k in vs[0]}
 z=dict(m.get("fixed")or{});vals=d.get("fields")or{};F=m.get("fields")or{};anchors={}
 warn+=[f"diapo {d['_i']} ({n}) : champ inconnu « {k} » ignoré"for k in vals if k not in F]
 for k,f in F.items():
  v=vals.get(k,f.get("default"))
  if v in(None,"",[]):z.setdefault(f["zone"],None);continue  # champ non fourni -> zone supprimée
  if f.get("kind")=="table":
   if len(v)>f.get("max_rows",99)or max(map(len,v))>f.get("max_cols",99):warn.append(f"diapo {d['_i']} ({n}) : « {k} » dépasse {f.get('max_rows')} lignes x {f.get('max_cols')} colonnes")
  else:
   L=[str(x)for x in v]if isinstance(v,list)else str(v).split("\n")
   if f.get("max_lines")and len(L)>f["max_lines"]:warn.append(f"diapo {d['_i']} ({n}) : « {k} » a {len(L)} lignes (max {f['max_lines']})")
   for x in L:
    if f.get("max")and len(plain(x))>f["max"]:warn.append(f"diapo {d['_i']} ({n}) : « {k} » trop long ({len(plain(x))} car., max {f['max']}) : {plain(x)[:40]}…")
   u=f.get("upper");u=len(L)if u is True else int(u or 0);L=[x.upper()if i<u else x for i,x in enumerate(L)]  # majuscules saisies dans le modèle
   if f.get("breaks"):L=["\v".join(L)]  # lignes = sauts de ligne d'un même paragraphe
   v=(f.get("prefix")or[])+L
  z[f["zone"]]=v
  if f.get("anchor")or f.get("bullets")is False:anchors[f["zone"]]={"anchor":f.get("anchor"),"bullets":f.get("bullets",True)}
 return[{**src,"zones":z,"anchors":anchors,"notes":d.get("notes")}]
def make_pptx(title,slides,subtitle="",template_prefix="",warn=None,ambiance="",uid=""):
 c=catalog(template_prefix);U=_umodels(uid,template_prefix)if uid else{}
 if U:c={**c,"models":{**c["models"],**U}}  # diapos de l'utilisateur (catalogue étendu)
 tn=c["template"];p=Presentation(_untemplate(c["bytes"]));orig=list(p.slides);uprs={}
 if not orig:raise ValueError("Le modèle PowerPoint ne contient aucune diapo type")
 p.core_properties.title=title;p.core_properties.language=LANG;fonts=(c["fonts"].get("accent",""),c["fonts"].get("texte",""));slides=list(slides);warn=[]if warn is None else warn;used={}
 amb=ambiance or c["default_ambiance"]
 if amb and c["ambiances"]and amb not in c["ambiances"]:warn.append(f"ambiance « {amb} » inconnue ({', '.join(c['ambiances'])})")
 if any("model"in d for d in slides):cover=[];slides=[x for i,d in enumerate(slides,1)for x in(_expand(c,{**d,"_i":i},used,warn,amb)if"model"in d else[d])]
 else:cover=[]if slides and str(slides[0].get("type"))=="1"else[{"type":1,"title":title,"bullets":[subtitle]if subtitle else[]}]
 for d in cover+slides:
  tb=d.get("table");body=[d.get("bullets")or[],d.get("bullets2")or[]];bi=0;done=set()
  if d.get("user"):
   if d["user"]not in uprs:uprs[d["user"]]=Presentation(str(_udir(uid,template_prefix)/d["user"]))
   s=_import(p,uprs[d["user"]].slides[int(d["slide"])-1])
  else:s=p.slides.add_slide(_layout(p,d["layout"]))if d.get("layout")else _dup(p,_pick(p,d.get("type"),bool(tb)))
  for k,v,shs in[(k,v,_zone(s,k))for k,v in(d.get("zones")or{}).items()]:  # résolution avant suppression : « Nom#n » ne se décale pas
   for sh in shs:
    done.add(sh.shape_id)
    if v is None:sh._element.getparent().remove(sh._element)  # null -> forme supprimée (ex. étiquette « EXEMPLE »)
    elif sh.has_table and isinstance(v,list):_table(sh.table,v,fonts)
    elif sh.has_text_frame:
     _fill(sh.text_frame,v if isinstance(v,list)else str(v).split("\n"),fonts)
     fm=(d.get("anchors")or{}).get(k)or{}
     if fm.get("anchor"):sh.text_frame._txBody.bodyPr.set("anchor",fm["anchor"])
     if fm.get("bullets")is False:  # texte sans puce dans une zone à puces héritées
      for pa in sh.text_frame.paragraphs:
       pr=pa._p.get_or_add_pPr();pr.set("marL","0");pr.set("indent","0")
       for t in("a:buChar","a:buAutoNum","a:buBlip","a:buNone"):
        for x in pr.findall(pq(t)):pr.remove(x)
       nx=next((x for x in pr if x.tag in(pq("a:tabLst"),pq("a:defRPr"),pq("a:extLst"))),None);bn=pr.makeelement(pq("a:buNone"),{})
       nx.addprevious(bn)if nx is not None else pr.append(bn)  # ordre du schéma DrawingML
  for sh in[]if d.get("asis")or d.get("user")else list(s.shapes):  # asis : diapo insérée telle quelle ; diapo utilisateur : hors champs = décor conservé
   if sh.shape_id in done:continue
   ph=sh.is_placeholder and int(sh.placeholder_format.type)
   if sh.has_table and tb:_table(sh.table,tb,fonts);tb=None
   elif ph in(1,3):_fill(sh.text_frame,[d.get("title","")],fonts)
   elif ph in(2,4,7)and sh.has_text_frame:
    if bi<2 and body[bi]:_fill(sh.text_frame,body[bi],fonts);bi+=1
    else:sh._element.getparent().remove(sh._element)  # zone vide -> supprimée (pas de texte d'exemple)
   elif ph==18 and sh._element.tag==pq("p:sp"):sh._element.getparent().remove(sh._element)  # espace image vide
  if tb:W,H=p.slide_width,p.slide_height;_table(s.shapes.add_table(len(tb),max(map(len,tb)),Emu(W//20),Emu(H//4),Emu(W*9//10),Emu(H//2)).table,tb,fonts)
  if d.get("notes"):s.notes_slide.notes_text_frame.text=d["notes"]
 sl=p.slides._sldIdLst
 for s in orig:
  for x in list(sl):
   if p.part.related_part(x.rId)is s.part:p.part.drop_rel(x.rId);sl.remove(x)
 return tn,p

# ---------- Diapos utilisateur : extension du catalogue, avec contrôle d'écart au modèle d'entreprise ----------
# Un utilisateur dépose une présentation qu'il a faite (lien de dépôt), chaque diapo est notée (0-100) par rapport au modèle épinglé,
# les diapos retenues deviennent des « modèles » de son catalogue : le chat en remplit les textes, décor et images sont conservés.
IMP=CAT/"imports";MAXUP=int(E("MAX_UPLOAD","52428800"))
SEUILS={"alerte":80,"refus":60}
def _uid(ctx):
 """Utilisateur LibreChat (en-tête X-User-Id, cf. {{LIBRECHAT_USER_ID}}) ; anonymisé ; « commun » à défaut."""
 try:v=ctx.request_context.request.headers.get("x-user-id","")
 except Exception:v=""
 return hashlib.sha256(v.encode()).hexdigest()[:16]if v.strip()else"commun"
def _udir(uid,prefix=""):return CAT/"users"/uid/(re.sub(r"\W","_",prefix)or"pptx")
def _umodels(uid,prefix=""):
 f=_udir(uid,prefix)/"models.json"
 if not f.is_file():return{}
 M=json.loads(f.read_text("utf8"))
 for m in M.values():
  m["_v"]=[{"user":m["source"]["user"],"slide":m["source"]["slide"]}]
  for x in(m.get("fields")or{}).values():x["_user"]=True
 return M
def _palette(c):
 """Couleurs de la charte : thème + toutes les couleurs employées dans le modèle (masques, dispositions, diapos types)."""
 if"palette"not in c:
  z=zipfile.ZipFile(io.BytesIO(c["bytes"]));s=set()
  for n in z.namelist():
   if n.endswith(".xml")and n.startswith(("ppt/theme","ppt/slideMasters","ppt/slideLayouts","ppt/slides/")):s|={x.upper()for x in re.findall(r'(?:srgbClr val|lastClr)="([0-9A-Fa-f]{6})"',z.read(n).decode("utf8","ignore"))}
  c["palette"]=sorted(s)
 return c["palette"]
def _near(h,pal,tol=28):
 r,g,b=(int(h[i:i+2],16)for i in(0,2,4))
 if max(r,g,b)-min(r,g,b)<=12:return True  # gris, noir, blanc : neutres
 return any(abs(r-int(x[0:2],16))+abs(g-int(x[2:4],16))+abs(b-int(x[4:6],16))<=tol for x in pal)
def _theme(master):
 b=master.part.part_related_by(RT.THEME).blob.decode("utf8","ignore")
 fo=tuple((m[1]if m else"")for m in(re.search(rf"<a:{k}Font>\s*<a:latin typeface=\"([^\"]*)\"",b)for k in("major","minor")))
 cs=re.search(r"<a:clrScheme.*?</a:clrScheme>",b,re.S);cl=dict(re.findall(r'<a:(dk1|lt1|dk2|lt2|accent\d|hlink|folHlink)>\s*<a:(?:srgbClr val|sysClr[^>]*lastClr)="([0-9A-Fa-f]{6})"',cs[0]if cs else""))
 return fo,{k:v.upper()for k,v in cl.items()}
def _sig(lay):return sorted((h.placeholder_format.idx,int(h.placeholder_format.type))for h in lay.placeholders)
def _fields(s):
 """Champs proposés pour une diapo utilisateur : chaque zone de texte ou tableau, dans l'ordre de lecture."""
 names=[sh.name for sh in s.shapes];seen={};F={};k=0;items=[]
 for sh in s.shapes:
  seen[sh.name]=seen.get(sh.name,0)+1;key=sh.name if names.count(sh.name)==1 else f"{sh.name}#{seen[sh.name]}"
  if sh.has_table or(sh.has_text_frame and sh.text_frame.text.strip()and not re.fullmatch(r"\s*[#\d]{1,2}([.,]\d{1,2})?\s*",sh.text_frame.text)):items.append((sh.top or 0,sh.left or 0,key,sh))
 for _,_,key,sh in sorted(items,key=lambda x:(round(x[0]/18e4),x[1])):
  if sh.has_table:F[f"tableau{sum(1 for x in F if x.startswith('tableau'))+1}"]={"zone":key,"kind":"table","max_rows":len(sh.table.rows)+3,"max_cols":len(sh.table.columns),"exemple":" | ".join(c.text for c in sh.table.rows[0].cells)[:80]};continue
  ps=[p.text.replace("\v"," ")for p in sh.text_frame.paragraphs if p.text.strip()]
  t=sh.is_placeholder and int(sh.placeholder_format.type)in(1,3)and"titre"not in F
  if not t:k+=1
  F["titre"if t else f"texte{k}"]={"zone":key,**({"kind":"lines","max_lines":len(ps)+2}if len(ps)>1 else{}),"max":max(25 if t else 20,round(max(map(len,ps))*1.3)),"exemple":" / ".join(ps)[:120]}
 return F
def conformity(c,tpl,up,n):
 """Écart d'une diapo utilisateur au modèle d'entreprise : score 0-100, niveau, alertes, champs proposés.
 Toute entorse à la charte (polices, couleurs, thème) donne au moins une alerte ; contenu non reproductible ou disposition inconnue : refus."""
 s=up.slides[n-1];a=[];pen=0;bloc=False;ch=False
 def add(p,msg,b=False,charte=False):
  nonlocal pen,bloc,ch;pen+=p;bloc|=b;a.append(msg);ch|=charte
 if(up.slide_width,up.slide_height)!=(tpl.slide_width,tpl.slide_height):add(20,f"format de diapo différent ({up.slide_width/36e4:.1f} x {up.slide_height/36e4:.1f} cm)")
 L=next((l for m in tpl.slide_masters for l in m.slide_layouts if l.name==s.slide_layout.name),None)
 if L is None:add(50,f"disposition « {s.slide_layout.name} » absente du modèle d'entreprise",True)
 elif _sig(L)!=_sig(s.slide_layout):add(10,f"disposition « {L.name} » modifiée par rapport au modèle")
 fo,cl=_theme(s.slide_layout.slide_master);tfo,tcl=_theme((L or tpl.slide_layouts[0]).slide_master)
 if fo!=tfo:add(25,charte=True,msg=f"polices du thème différentes ({' / '.join(fo)} au lieu de {' / '.join(tfo)})")
 dc=[k for k in tcl if cl.get(k)!=tcl[k]]
 if dc:add(min(20,3*len(dc)),charte=True,msg=f"couleurs du thème différentes ({', '.join(dc)})")
 x=s._element.xml;fam=(c["fonts"].get("texte")or"").split(" ")[0]
 ff=sorted({f for f in re.findall(r'<a:latin typeface="([^"]+)"',x)if not f.startswith(("+",fam))})
 if ff:add(min(30,15*len(ff)),charte=True,msg=f"polices hors charte : {', '.join(ff)}")
 pal=_palette(c);oc=sorted({h.upper()for h in re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"',x)if not _near(h,pal)})
 if oc:add(min(30,10*len(oc)),charte=True,msg=f"couleurs hors palette : {', '.join('#'+h for h in oc[:6])}")
 if any(int(v)<800 for v in re.findall(r'<a:rPr[^>]*\ssz="(\d+)"',x)):add(5,"texte de moins de 8 pt (lisibilité)")
 W,H=up.slide_width,up.slide_height;out=[sh.name for sh in s.shapes if sh.left is not None and(sh.left<-W*.02 or sh.top<-H*.02 or sh.left+sh.width>W*1.02 or sh.top+sh.height>H*1.02)]
 if out:add(min(15,5*len(out)),f"formes hors de la diapo : {', '.join(out[:4])}")
 rt=[r.reltype.rsplit("/",1)[-1]for r in s.part.rels._rels.values()]
 un=sorted({t for t in rt if t in("chart","diagramData","diagramLayout","oleObject","package","video","audio","media")})
 if un:add(50,f"contenu non reproductible : {', '.join(un)} (graphique, SmartArt, objet incorporé ou média)",True)
 F=_fields(s)
 if not F:a.append("aucune zone de texte : la diapo sera insérée telle quelle")
 sc=max(0,100-pen);S={**SEUILS,**(c.get("conformite")or{})}
 lvl="refus"if bloc or sc<S["refus"]else"alerte"if sc<S["alerte"]or ch else"conforme"
 return{"diapo":n,"titre":(s.shapes.title.text_frame.text.strip()[:60]if s.shapes.title is not None and s.shapes.title.has_text_frame else"")or s.slide_layout.name,"score":sc,"niveau":lvl,"alertes":a,"champs":F}
def analyze_upload(tok,prefix=""):
 c=catalog(prefix);d=IMP/tok;up=Presentation(str(d/"deck.pptx"));tpl=Presentation(_untemplate(c["bytes"]))
 r={"fichier":json.loads((d/"meta.json").read_text("utf8")).get("name",""),"modele":c["template"],"seuils":{**SEUILS,**(c.get("conformite")or{})},"diapos":[conformity(c,tpl,up,i)for i in range(1,len(up.slides)+1)]}
 (d/"report.json").write_text(json.dumps(r,ensure_ascii=False,indent=1),"utf8");return r
def _report_txt(r):
 ic={"conforme":"✓","alerte":"⚠","refus":"✗"}
 return(f"Analyse de « {r['fichier']} » par rapport au modèle « {r['modele']} » (conforme ≥ {r['seuils']['alerte']}, refus < {r['seuils']['refus']}) :\n"
  +"\n".join(f"{ic[x['niveau']]} diapo {x['diapo']} « {x['titre']} » : {x['score']}/100 ({x['niveau']})"+"".join(f"\n    - {m}"for m in x["alertes"])
   +("\n    champs : "+"; ".join(f"{k} ({v.get('kind','texte')}, ex. {v['exemple'][:40]!r})"for k,v in x["champs"].items())if x["champs"]else"")for x in r["diapos"]))
UPLOAD_PAGE="""<!doctype html><html lang="fr"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Ajouter mes diapos</title>
<style>body{font:15px/1.5 system-ui,sans-serif;max-width:46rem;margin:2rem auto;padding:0 1rem;color:#1d1a24}h1{font-size:1.4rem}
.zone{border:2px dashed #b9a6e8;border-radius:10px;padding:1.5rem;text-align:center;background:#f6f2fd}button{font:inherit;padding:.5rem 1rem;border-radius:6px;border:0;background:#6d32d9;color:#fff;cursor:pointer}
li{margin:.6rem 0;padding:.6rem .8rem;border-radius:8px;list-style:none}.conforme{background:#e3f3e9}.alerte{background:#fbefdc}.refus{background:#fbe2e2}ul{padding:0}small{color:#5c566b}</style>
<h1>Ajouter des diapos à mon catalogue</h1><p>Déposez une présentation réalisée avec le modèle d'entreprise. Chaque diapo est comparée au modèle : <b>conforme</b>, <b>alerte</b> (écart notable) ou <b>refus</b> (trop éloignée). Revenez ensuite dans la conversation pour choisir les diapos à ajouter.</p>
<div class="zone"><input type="file" id="f" accept=".pptx"> <button id="b">Analyser</button><p id="s"></p></div><ul id="r"></ul>
<script>
const s=document.getElementById('s'),r=document.getElementById('r');
document.getElementById('b').onclick=async()=>{const f=document.getElementById('f').files[0];if(!f){s.textContent='Choisissez un fichier .pptx.';return}
s.textContent='Analyse en cours…';r.innerHTML='';
const x=await fetch(location.pathname+'?name='+encodeURIComponent(f.name),{method:'PUT',body:f});const j=await x.json();
if(!x.ok){s.textContent=j.error||'Erreur';return}
s.textContent=`${j.diapos.length} diapos analysées (modèle ${j.modele}). Indiquez dans la conversation les numéros à ajouter.`;
for(const d of j.diapos){const li=document.createElement('li');li.className=d.niveau;
li.innerHTML=`<b>Diapo ${d.diapo} · ${d.score}/100 · ${d.niveau}</b> — ${d.titre.replace(/</g,'&lt;')}<br><small>${(d.alertes.join(' ; ')||'aucun écart détecté').replace(/</g,'&lt;')}</small>`;r.appendChild(li)}};
</script></html>"""

# ---------- Serveur MCP ----------
mcp=FastMCP("office-templates",host=E("HOST","0.0.0.0"),port=int(E("PORT","8000")),stateless_http=True)
def _chk(*a):
 if len(repr(a))>MAXIN:raise ValueError("Contenu trop volumineux")
def _ret(kind,name,tn,obj):u=_save(kind,name,obj);return f"Document généré avec le modèle « {tn} » : [{name}.{kind}]({u}) (lien valable {TTL//60} min)"

@mcp.tool()
def list_templates()->str:
 """Liste les modèles d'entreprise disponibles et indique le plus récent par type."""
 f=_list();r=[]
 for k,e in EXT.items():
  c=sorted([x for x in f if x["name"].lower().endswith(e)],key=lambda x:x["modified"],reverse=True)
  r.append(f"{k}: "+(", ".join(x["name"]for x in c)or"aucun")+(f" (utilisé: {c[0]['name']})"if c else""))
 return"\n".join(r)
@mcp.tool()
def create_word(title:str,markdown:str,filename:str="",template_prefix:str="")->str:
 """Crée un document Word (.docx) avec le dernier modèle d'entreprise. `markdown`: contenu (titres #, listes -, 1., tableaux |, **gras**, *italique*, blocs ```). `template_prefix`: filtre optionnel sur le nom du modèle."""
 _chk(title,markdown);tn,d=make_docx(title,markdown,template_prefix);return _ret("docx",filename or title,tn,d)
@mcp.tool()
def create_excel(title:str,sheets:list[dict],filename:str="",template_prefix:str="")->str:
 """Crée un classeur Excel (.xlsx) avec le dernier modèle. `sheets`: [{"name":"Ventes","rows":[["Col1","Col2"],[1,2]]}], 1re ligne = en-têtes (mise en tableau Excel)."""
 _chk(title,sheets);tn,w=make_xlsx(title,sheets,template_prefix);return _ret("xlsx",filename or title,tn,w)
@mcp.tool()
def get_presentation_catalog(ctx:Context,template_prefix:str="")->str:
 """À appeler avant create_powerpoint : diapos disponibles dans le modèle d'entreprise (usage, champs à remplir, limites) et règles de rédaction, plus les diapos ajoutées par l'utilisateur (« Mes diapos »)."""
 c=catalog(template_prefix);f=c["fonts"];U=_umodels(_uid(ctx),template_prefix)
 if U:c={**c,"models":{**c["models"],**U},"groups":(c["groups"]or[{"name":"Diapos","models":list(c["models"])}])+[{"name":"Mes diapos","help":"ajoutées par l'utilisateur depuis ses présentations ; texte remplaçable, décor et images conservés","models":list(U)}]}
 def fd(k,x):
  lim=", ".join(filter(None,[x.get("kind","text")if x.get("kind")in("lines","table")else"texte",f"≤{x['max']} car."if x.get("max")else"",f"≤{x['max_lines']} lignes"if x.get("max_lines")else"",f"≤{x.get('max_rows')}x{x.get('max_cols')}"if x.get("kind")=="table"else""]))
  return f"{k} ({lim})"+(f" : {x['help']}"if x.get("help")else"")+(f" ex. {x['exemple'][:50]!r}"if x.get("_user")and x.get("exemple")else"")+(f" [défaut : {x['default']!r}]"if x.get("default")else"")
 def var(m):
  V=m["_v"]
  if m.get("sequence"):return f" [{len(V)} diapos insérées]"
  if m.get("variant_by"):return f" [variante choisie par « {m['variant_by']} »]"
  if len(V)<2:return""
  a=sorted({x["ambiance"]for x in V if x.get("ambiance")})
  return f" [{len(V)} variantes"+(f", ambiances {'/'.join(a)}"if a else"")+" ; omis = alternance automatique, « variante »: n pour imposer : "+", ".join(f"{i}={x.get('label',next(iter(x.values())))}"for i,x in enumerate(V,1))+"]"
 def model(n,m):
  F=m.get("fields")or{}
  cf=m.get("conformite")or{};al=f" ⚠ ÉCART AU MODÈLE ({cf['score']}/100 : {'; '.join(cf['alertes'][:3])}) — à signaler à l'utilisateur"if cf.get("niveau")in("alerte","refus")else""
  return f"* {n} — {m.get('usage','')}{var(m)}{al}\n  "+("(contenu fixe, aucun champ)"if m.get("asis")else"champs : "+"; ".join(fd(k,x)for k,x in F.items()))
 G=c["groups"]or[{"name":"Diapos","models":list(c["models"])}];seen=set(x for g in G for x in g["models"])
 G=G+([{"name":"Autres","models":[n for n in c["models"]if n not in seen]}]if any(n not in seen for n in c["models"])else[])
 return(f"Modèle « {c['template']} » (analysé le {c['analyzed']}). Polices appliquées automatiquement : {f.get('texte')} (texte), {f.get('accent')} (accent).\n"
  +("Règles :\n"+"\n".join(f"- {r}"for r in c["rules"])+"\n"if c["rules"]else"")
  +("Ambiances couleur (paramètre « ambiance », une par document) : "+"; ".join(f"{k} = {v}"for k,v in c["ambiances"].items())+(f" (défaut : {c['default_ambiance']})"if c["default_ambiance"]else"")+"\n"if c["ambiances"]else"")
  +"\n".join(f"\n## {g['name']}"+(f" — {g['help']}"if g.get("help")else"")+"\n"+"\n".join(model(n,c["models"][n])for n in g["models"]if n in c["models"])for g in G)
  +'\n\nAppel : create_powerpoint(title, slides=[{"model":"…","fields":{"champ":"texte" | ["ligne",…] | [["cellule",…],…]},"variante":n (facultatif),"notes":"…"}], ambiance="…"). Champ omis = zone retirée.')
@mcp.tool()
def refresh_template(template_prefix:str="")->str:
 """À n'appeler que sur demande explicite (nouveau modèle publié) : épingle le modèle PowerPoint le plus récent, refait l'analyse (polices, zones) et vérifie la configuration des diapos."""
 c=catalog(template_prefix,refresh=True);inv=c["inventory"]
 return(f"Modèle épinglé : {c['template']} ; polices {c['fonts']} ; {len(inv['slides'])} diapos types, {len(inv['layouts'])} dispositions dont {sum(map(len,inv['families'].values()))} en {len(inv['families'])} familles de variantes ; {len(c['models'])} modèles de diapos exposés."
  +("\n⚠ "+"\n⚠ ".join(c["warnings"])if c["warnings"]else" Configuration cohérente."))
@mcp.tool()
def list_slide_types(template_prefix:str="",layouts:bool=False)->str:
 """(Administration) Inventaire brut du modèle épinglé : diapos types et zones, `layouts=True` pour les dispositions (zones « @idx »). Sert à rédiger catalog/pptx.json."""
 c=catalog(template_prefix);inv=c["inventory"]
 r=f"Modèle {c['template']}\n"+"\n".join(f"{t['type']}. {t['label']} | zones: "+"; ".join(f"{z['zone']} ({z['kind']}: {z['exemple'][:40]!r})"for z in t["zones"])for t in inv["slides"])
 r+="\n\nFamilles de variantes (source « family ») :\n"+"\n".join(f"- {n} : {', '.join(v)}"for n,v in inv["families"].items())
 if layouts:r+="\n\nDispositions (source « layout », masque n) :\n"+"\n".join(f"- {l['layout']} (masque {l.get('master',1)}) | "+("; ".join(f"{z['zone']} ({z['kind']}: {z['exemple'][:30]!r})"for z in l["zones"])or"mise en page figée")for l in inv["layouts"])
 return r
@mcp.tool()
def create_powerpoint(ctx:Context,title:str,slides:list[dict],ambiance:str="",filename:str="",template_prefix:str="")->str:
 """Crée une présentation (.pptx) strictement conforme au modèle d'entreprise. Appeler d'abord get_presentation_catalog.
 `slides`: [{"model":"couverture","fields":{"titre":"…"}},…] dans l'ordre voulu ; "variante": n impose une variante visuelle ; **mot** = mise en valeur (police d'accent) ; « \\n » dans un texte = nouveau paragraphe ; "notes" = notes orateur.
 `ambiance` : couleur dominante du document (cf. catalogue) ; les variantes (couverture, intercalaires…) sont choisies dans cette ambiance et alternées.
 Les dépassements de limites sont signalés : corriger le contenu et regénérer si besoin."""
 _chk(title,slides);w=[];tn,p=make_pptx(title,slides,template_prefix=template_prefix,warn=w,ambiance=ambiance,uid=_uid(ctx))
 return _ret("pptx",filename or title,tn,p)+("\n\nÀ corriger :\n- "+"\n- ".join(w)if w else"")

@mcp.tool()
def add_slides_link(ctx:Context,template_prefix:str="")->str:
 """Étendre le catalogue avec des diapos de l'utilisateur : renvoie un lien de dépôt (valable 24 h) où il dépose une présentation .pptx faite avec le modèle d'entreprise ; chaque diapo y est comparée au modèle. Ensuite : get_slides_report puis add_user_slides."""
 IMP.mkdir(parents=True,exist_ok=True)
 for d in IMP.iterdir():
  if d.stat().st_mtime<time.time()-86400:shutil.rmtree(d,ignore_errors=True)
 tok=secrets.token_urlsafe(24);(IMP/tok).mkdir();(IMP/tok/"meta.json").write_text(json.dumps({"uid":_uid(ctx),"prefix":template_prefix,"created":time.time()}),"utf8")
 return f"Lien de dépôt (valable 24 h) : {BASE}/import/{tok}\nIdentifiant d'import : {tok}\nAprès le dépôt, appeler get_slides_report(\"{tok}\")."
def _imp(tok,ctx):
 d=IMP/str(tok)
 if not re.fullmatch(r"[\w-]{20,64}",str(tok))or not(d/"meta.json").is_file():raise ValueError("Import inconnu ou expiré : redemander un lien avec add_slides_link")
 m=json.loads((d/"meta.json").read_text("utf8"))
 if m["uid"]!=_uid(ctx):raise ValueError("Cet import appartient à un autre utilisateur")
 if not(d/"deck.pptx").is_file():raise ValueError("Aucun fichier déposé pour l'instant : l'utilisateur doit ouvrir le lien et déposer sa présentation")
 return d,m
def _rep(d,tok,m):return json.loads((d/"report.json").read_text("utf8"))if(d/"report.json").is_file()else analyze_upload(tok,m["prefix"])
@mcp.tool()
def get_slides_report(ctx:Context,import_id:str)->str:
 """Rapport d'analyse d'une présentation déposée : pour chaque diapo, score d'écart au modèle d'entreprise (0-100), niveau (conforme / alerte / refus), écarts détectés et champs qui seront proposés. Présenter les alertes à l'utilisateur."""
 d,m=_imp(import_id,ctx)
 return _report_txt(_rep(d,import_id,m))+"\n\nPour ajouter : add_user_slides(import_id, slides=[{\"diapo\":n,\"nom\":\"mon_modele\",\"usage\":\"à quoi sert la diapo\"}]). Diapos en « refus » : seulement avec forcer=true, après accord explicite de l'utilisateur."
@mcp.tool()
def add_user_slides(ctx:Context,import_id:str,slides:list[dict],forcer:bool=False)->str:
 """Ajoute des diapos analysées au catalogue de l'utilisateur. `slides`: [{"diapo":3,"nom":"bilan_projet","usage":"…","champs":{"texte1":"col1_titre",…}}]. Renommer les champs d'après leurs exemples (rapport) pour que les noms disent leur rôle. Les diapos en alerte sont ajoutées avec leur alerte, à signaler à l'utilisateur ; en refus, seulement si forcer=true."""
 d,m=_imp(import_id,ctx);r=_rep(d,import_id,m)
 base=catalog(m["prefix"])["models"];ud=_udir(m["uid"],m["prefix"]);ud.mkdir(parents=True,exist_ok=True)
 b=(d/"deck.pptx").read_bytes();fn=hashlib.sha256(b).hexdigest()[:16]+".pptx";(ud/fn).write_bytes(b)
 mf=ud/"models.json";M=json.loads(mf.read_text("utf8"))if mf.is_file()else{};out=[]
 for x in slides:
  n=int(x.get("diapo",0));nom=str(x.get("nom","")).strip().lower()
  a=next((y for y in r["diapos"]if y["diapo"]==n),None)
  if not a:out.append(f"✗ diapo {n} : absente du fichier");continue
  if not re.fullmatch(r"[a-z0-9_]{3,40}",nom)or nom in base:out.append(f"✗ diapo {n} : nom « {nom} » invalide ou déjà utilisé par le modèle d'entreprise (a-z, 0-9, _)");continue
  if a["niveau"]=="refus"and not forcer:out.append(f"✗ diapo {n} : refusée, trop éloignée du modèle ({a['score']}/100) : {'; '.join(a['alertes'])}");continue
  F={(x.get("champs")or{}).get(k,k):v for k,v in a["champs"].items()}
  M[nom]={"usage":str(x.get("usage")or a["titre"])[:200],"source":{"user":fn,"slide":n},"fields":F,"origine":f"{r['fichier']}, diapo {n}","ajoute":time.strftime("%Y-%m-%d"),
   "conformite":{"score":a["score"],"niveau":a["niveau"],"alertes":a["alertes"],"force":a["niveau"]=="refus"}}
  out.append(("⚠"if a["niveau"]!="conforme"else"✓")+f" « {nom} » ajouté ({a['score']}/100, {a['niveau']})"+(f" — écarts : {'; '.join(a['alertes'])}"if a["niveau"]!="conforme"else""))
 mf.write_text(json.dumps(M,ensure_ascii=False,indent=1),"utf8")
 return"\n".join(out)+"\nCes diapos apparaissent dans get_presentation_catalog (rubrique « Mes diapos »)."
@mcp.tool()
def remove_user_slide(ctx:Context,nom:str,template_prefix:str="")->str:
 """Retire une diapo du catalogue personnel de l'utilisateur (rubrique « Mes diapos »)."""
 mf=_udir(_uid(ctx),template_prefix)/"models.json";M=json.loads(mf.read_text("utf8"))if mf.is_file()else{}
 if nom not in M:return f"« {nom} » ne fait pas partie de vos diapos ({', '.join(M)or'aucune'})"
 del M[nom];mf.write_text(json.dumps(M,ensure_ascii=False,indent=1),"utf8");return f"« {nom} » retiré de vos diapos."

@mcp.custom_route("/import/{tok}",methods=["GET","PUT"])
async def upload(req):
 from starlette.concurrency import run_in_threadpool
 tok=req.path_params["tok"];d=IMP/tok
 if not re.fullmatch(r"[\w-]{20,64}",tok)or not(d/"meta.json").is_file()or(d/"meta.json").stat().st_mtime<time.time()-86400:return HTMLResponse("<p>Lien de dépôt inconnu ou expiré. Redemandez un lien dans la conversation.</p>",404)
 if req.method=="GET":return HTMLResponse(UPLOAD_PAGE,headers={"Cache-Control":"no-store","X-Content-Type-Options":"nosniff"})
 if int(req.headers.get("content-length")or 0)>MAXUP:return JSONResponse({"error":f"Fichier trop volumineux (max {MAXUP//1048576} Mo)"},413)
 b=await req.body()
 if len(b)>MAXUP:return JSONResponse({"error":"Fichier trop volumineux"},413)
 try:
  if"ppt/presentation.xml"not in zipfile.ZipFile(io.BytesIO(b)).namelist():raise ValueError
 except Exception:return JSONResponse({"error":"Ce fichier n'est pas une présentation PowerPoint (.pptx)"},400)
 m=json.loads((d/"meta.json").read_text("utf8"));m["name"]=re.sub(r"[^\w\-. ]","_",req.query_params.get("name","presentation.pptx"))[:120]
 (d/"deck.pptx").write_bytes(b);(d/"meta.json").write_text(json.dumps(m),"utf8");(d/"report.json").unlink(missing_ok=True)
 try:r=await run_in_threadpool(analyze_upload,tok,m.get("prefix",""))
 except Exception as e:return JSONResponse({"error":f"Analyse impossible : {e}"},400)
 return JSONResponse(r)
@mcp.custom_route("/files/{tok}/{name}",methods=["GET"])
async def dl(req):
 tok,name=req.path_params["tok"],req.path_params["name"]
 if not re.fullmatch(r"[\w-]{20,64}",tok)or"/"in name or".."in name:return Response(status_code=404)
 f=OUT/tok/name;kind=f.suffix[1:]
 if not f.is_file()or kind not in MIME or f.stat().st_mtime<time.time()-TTL:return Response(status_code=404)
 return Response(f.read_bytes(),media_type=MIME[kind],headers={"Content-Disposition":f"attachment; filename*=UTF-8''{quote(name)}","X-Content-Type-Options":"nosniff","Cache-Control":"no-store"})
@mcp.custom_route("/health",methods=["GET"])
async def health(_):return JSONResponse({"ok":True})

def app():
 a=mcp.streamable_http_app()
 async def auth(scope,rcv,snd):
  if scope["type"]=="http"and scope["path"].startswith("/mcp")and API_KEY:
   h=dict(scope["headers"]).get(b"authorization",b"").decode()
   if not hmac.compare_digest(h,"Bearer "+API_KEY):return await JSONResponse({"error":"unauthorized"},401)(scope,rcv,snd)
  await a(scope,rcv,snd)
 return auth
if __name__=="__main__":
 import sys
 if sys.argv[1:2]==["bundle"]:  # prépare catalog/bundle (modèle + analyse) à embarquer dans l'image, depuis TEMPLATE_DIR ou SharePoint
  CAT=BUNDLE;print(refresh_template(sys.argv[2]if len(sys.argv)>2 else""));sys.exit()
 if not API_KEY:print("⚠ MCP_API_KEY non défini : endpoint /mcp non protégé")
 try:c=catalog();print(f"PowerPoint : modèle « {c['template']} » (analysé le {c['analyzed']}), {len(c['models'])} diapos"+(f", {len(c['warnings'])} avertissement(s)"if c["warnings"]else""))
 except Exception as e:print(f"⚠ PowerPoint : aucun modèle disponible ({e}) ; fournir catalog/bundle, TEMPLATE_DIR ou SharePoint")
 uvicorn.run(app(),host=E("HOST","0.0.0.0"),port=int(E("PORT","8000")),proxy_headers=True)
