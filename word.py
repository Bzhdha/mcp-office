"""MCP Office – documents Word : catalogue de blocs, cadres (charte Niji ou contraintes client) et bibliothèque d'UO.
Le document part du modèle d'entreprise épinglé : page de garde, historique, interlocuteurs et sommaire sont remplis,
le corps d'exemple est remplacé par les blocs demandés, la présentation Niji et les CGV sont insérables, la 4e de couverture est conservée."""
import io,re,json,math,time,shutil,zipfile,subprocess,tempfile
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm,RGBColor,Emu
from docx.oxml import OxmlElement,parse_xml
from docx.oxml.ns import qn,nsdecls
from docx.enum.text import WD_BREAK

HERE=Path(__file__).with_name("catalog")
SYMB=("Symbol","Wingdings","Wingdings 2","Wingdings 3","Courier New","Segoe UI Symbol","MS Gothic")  # polices de puces : jamais remplacées
PAGES={"A4":(21.0,29.7),"A3":(29.7,42.0),"Letter":(21.59,27.94)}

# ---------- outils XML ----------
def _runs(t):return[(x.strip("*`"),x.startswith("**"),x.startswith("*")and not x.startswith("**"))for x in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)",str(t))if x]
def _shade(cell,fill):
 tcPr=cell._tc.get_or_add_tcPr()
 for s in tcPr.findall(qn("w:shd")):tcPr.remove(s)
 s=OxmlElement("w:shd");s.set(qn("w:val"),"clear");s.set(qn("w:color"),"auto");s.set(qn("w:fill"),fill);tcPr.append(s)
def _borders(tbl,color="D9D3E8",sz=4,inside=True):
 tblPr=tbl._tbl.tblPr;b=tblPr.find(qn("w:tblBorders"))
 if b is not None:tblPr.remove(b)
 b=OxmlElement("w:tblBorders")
 for e in("top","left","bottom","right")+(("insideH","insideV")if inside else()):
  x=OxmlElement(f"w:{e}");x.set(qn("w:val"),"single");x.set(qn("w:sz"),str(sz));x.set(qn("w:color"),color);b.append(x)
 tblPr.append(b)
def _cell_margins(tbl,top=60,left=100):
 tblPr=tbl._tbl.tblPr;m=OxmlElement("w:tblCellMar")
 for e,v in(("top",top),("left",left),("bottom",top),("right",left)):x=OxmlElement(f"w:{e}");x.set(qn("w:w"),str(v));x.set(qn("w:type"),"dxa");m.append(x)
 tblPr.append(m)
def _repeat_header(row):
 trPr=row._tr.get_or_add_trPr();h=OxmlElement("w:tblHeader");h.set(qn("w:val"),"true");trPr.append(h)
def _no_split(row):
 trPr=row._tr.get_or_add_trPr();c=OxmlElement("w:cantSplit");c.set(qn("w:val"),"true");trPr.append(c)
def _width(tbl,widths_cm):
 for row in tbl.rows:
  for c,w in zip(row.cells,widths_cm):c.width=Cm(w)
def _numpr(p,num,lvl):
 pPr=p._p.get_or_add_pPr();n=OxmlElement("w:numPr");a=OxmlElement("w:ilvl");a.set(qn("w:val"),str(lvl));b=OxmlElement("w:numId");b.set(qn("w:val"),str(num));n.append(a);n.append(b);pPr.insert(1 if pPr.find(qn("w:pStyle"))is not None else 0,n)

class Builder:
 """Construit un document à partir du modèle : cfg = catalog/docx.json, cadre = paramètres effectifs."""
 def __init__(self,tpl_bytes,cfg,cadre,amb):
  self.doc=Document(io.BytesIO(tpl_bytes));self.cfg=cfg;self.cad=cadre;self.warn=[]
  A=cfg["ambiances"];self.col=A.get(amb)or A[cfg.get("ambiance_defaut","violet")]
  if cadre.get("couleurs")=="sobre":self.col=A.get("sobre",{"fort":"404040","clair":"EDEDED","texte":"FFFFFF"})
  self.fonts=cadre.get("police_texte"),cadre.get("police_titres")
  self._styles();self._numbering();self.body=self.doc.element.body;self._marks()
 # --- styles de tableaux et listes à la charte ---
 def _styles(self):
  st=self.doc.styles.element;have={s.get(qn("w:styleId"))for s in st.findall(qn("w:style"))}
  lib=parse_xml((HERE/"docx-styles.xml").read_bytes())
  for s in lib.findall(qn("w:style")):
   if s.get(qn("w:styleId"))not in have:st.append(deepcopy(s))
 def _numbering(self):
  np=self.doc.part.numbering_part.element
  aid=max([int(x)for x in np.xpath("./w:abstractNum/@w:abstractNumId")]+[0])+1
  c=self.col["fort"]
  bul="".join(f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="{ch}"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="{357+340*i}" w:hanging="284"/></w:pPr><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial"/><w:color w:val="{c}"/><w:sz w:val="{16 if i==0 else 18}"/></w:rPr></w:lvl>'for i,ch in enumerate(["●","–","▪","–","▪","–","▪","–","▪"]))
  dec="".join(f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{"decimal"if i%3==0 else"lowerLetter"if i%3==1 else"lowerRoman"}"/><w:lvlText w:val="%{i+1}."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="{357+340*i}" w:hanging="357"/></w:pPr><w:rPr><w:color w:val="{c}"/><w:b/></w:rPr></w:lvl>'for i in range(9))
  last=np.findall(qn("w:abstractNum"))[-1]if np.findall(qn("w:abstractNum"))else None
  for k,body in((0,bul),(1,dec)):
   e=parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{aid+k}"><w:multiLevelType w:val="hybridMultilevel"/>{body}</w:abstractNum>')
   (last.addnext(e)if last is not None else np.insert(0,e));last=e
  self.np=np;self.aid=aid;self.bullet=self._num(aid)
 def _num(self,abstract,restart=False):
  nid=max([int(x)for x in self.np.xpath("./w:num/@w:numId")]+[0])+1
  e=parse_xml(f'<w:num {nsdecls("w")} w:numId="{nid}"><w:abstractNumId w:val="{abstract}"/>'+('<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride>'if restart else"")+"</w:num>")
  self.np.append(e);return nid
 # --- repères du modèle ---
 def _marks(self):
  R=self.cfg["reperes"];B=list(self.body);txt=lambda e:"".join(t.text or""for t in e.iter(qn("w:t")))
  sty=lambda e:(e.find(qn("w:pPr")+"/"+qn("w:pStyle")).get(qn("w:val"))if e.find(qn("w:pPr")+"/"+qn("w:pStyle"))is not None else"")
  h1=lambda e:e.tag==qn("w:p")and sty(e)in R["styles_titre1"]
  self.toc=next((e for e in B if e.tag==qn("w:sdt")),None);ti=B.index(self.toc)if self.toc is not None else 0
  self.i_body=next(i for i,e in enumerate(B)if i>ti and h1(e))
  self.i_niji=next((i for i,e in enumerate(B)if h1(e)and txt(e).strip().startswith(R["presentation"])),None)
  ic=next(i for i,e in enumerate(B)if R["cgv"]in txt(e))
  if ic>0 and B[ic-1].tag==qn("w:p")and not txt(B[ic-1]).strip()and"w:br"in B[ic-1].xml and'w:type="page"'in B[ic-1].xml:ic-=1
  self.i_cgv=ic;self.final=B[-1]
  self.niji=[deepcopy(e)for e in B[self.i_niji:ic]]if self.i_niji is not None else[]
  n=B[self.i_body].find(qn("w:pPr")+"/"+qn("w:numPr")+"/"+qn("w:numId"));self.hnum=n.get(qn("w:val"))if n is not None else None  # liste des titres numérotés
  self.tail=B[ic:-1];self.sect0=next(e for e in self.tail if e.find(".//"+qn("w:sectPr"))is not None)
  for e in B[self.i_body:-1]:self.body.remove(e)
 # --- page de garde, historique, interlocuteurs, sommaire ---
 def front(self,titre,sous_titre,historique,interlocuteurs,opt):
  R=self.cfg["reperes"];B=list(self.body)
  def put(p,text):
   rs=p.findall(".//"+qn("w:r"));ts=[r for r in rs if r.find(qn("w:t"))is not None]
   if not ts:return
   ts[0].find(qn("w:t")).text=text
   for r in ts[1:]:r.find(qn("w:t")).text=""
  for e in B:
   for p in e.iter(qn("w:p")):
    t="".join(x.text or""for x in p.iter(qn("w:t"))).strip()
    if t==R["titre_couverture"]:put(p,titre.upper())
    elif t==R["sous_titre_couverture"]:put(p,(sous_titre or"").upper())
    elif t.startswith(R.get("zone_projet","§"))and p.getparent().tag==qn("w:txbxContent"):put(p,titre)
    elif t.startswith(R.get("zone_document","§"))and p.getparent().tag==qn("w:txbxContent"):put(p,sous_titre or"")
  tbls=[e for e in self.body if e.tag==qn("w:tbl")]
  hist,inter=(tbls+[None,None])[:2]
  heads=[e for e in self.body if e.tag==qn("w:p")and"".join(x.text or""for x in e.iter(qn("w:t"))).strip()in(R["historique"],R["interlocuteurs"])]
  if hist is not None:
   rows=hist.findall(qn("w:tr"));tpl=rows[-1]
   for r in rows[1:]:hist.remove(r)
   for h in historique:
    r=deepcopy(tpl)
    for tc,v in zip(r.findall(qn("w:tc")),(h.get("version",""),h.get("date",""),h.get("modifications",""))):
     ts=list(tc.iter(qn("w:t")))
     if ts:ts[0].text=str(v);[setattr(x,"text","")for x in ts[1:]]
     else:tc.find(qn("w:p")).append(parse_xml(f'<w:r {nsdecls("w")}><w:t xml:space="preserve">{_esc(v)}</w:t></w:r>'))
    hist.append(r)
  if inter is not None:
   rows=inter.findall(qn("w:tr"))
   for r in rows[len(interlocuteurs):]:inter.remove(r)
   for r,c in zip(rows,interlocuteurs):
    for d in r.iter(qn("w:drawing")):d.getparent().remove(d)  # pas de photo d'exemple
    ps=[p for p in r.iter(qn("w:p"))if"".join(x.text or""for x in p.iter(qn("w:t"))).strip()]
    for p,v in zip(ps,(c.get("nom",""),c.get("fonction",""),c.get("email",""),c.get("telephone",""))):put(p,v)
    for p in ps[4:]:put(p,"")
  if not opt.get("historique",True)and hist is not None:
   self.body.remove(hist);[self.body.remove(h)for h in heads if"".join(x.text or""for x in h.iter(qn("w:t"))).strip()==R["historique"]]
  if(not interlocuteurs or not opt.get("interlocuteurs",True))and inter is not None:
   self.body.remove(inter);[self.body.remove(h)for h in heads if"".join(x.text or""for x in h.iter(qn("w:t"))).strip()==R["interlocuteurs"]]
  if not opt.get("sommaire",True)and self.toc is not None:
   nx=self.toc.getnext();self.body.remove(self.toc)
   if nx is not None and nx.tag==qn("w:p")and'w:type="page"'in nx.xml and not"".join(nx.itertext()).strip():self.body.remove(nx)
  if not opt.get("page_de_garde",True):
   for e in list(self.body)[:self.body.index(next(e for e in self.body if e.tag in(qn("w:tbl"),qn("w:sdt"))or(e.tag==qn("w:p")and"".join(x.text or""for x in e.iter(qn("w:t"))).strip()==R["historique"])))]:
    if e.find(".//"+qn("w:sectPr"))is None:self.body.remove(e)
  s=self.doc.settings.element;u=s.find(qn("w:updateFields"))
  if u is None:u=OxmlElement("w:updateFields");s.append(u)
  u.set(qn("w:val"),"true")  # le sommaire est recalculé à l'ouverture
 # --- texte ---
 def _para(self,box,text,style=None,lvl=None,num=None,bold=False,color=None,size=None,first=None):
  p=first if first is not None else box.add_paragraph(style=style)
  if first is not None and style:p.style=self.doc.styles[style]
  if num is not None:_numpr(p,num,lvl or 0)
  for t,b,i in _runs(text):
   r=p.add_run(t)
   if b or bold:
    if self.cad.get("gras_police")and not bold:r.font.name=self.cad["gras_police"]
    else:r.bold=True
   if i:r.italic=True
   if color:r.font.color.rgb=RGBColor.from_string(color)
   if size:r.font.size=Pt(size)
  return p
 def lines(self,box,lines,size=None,first=None):
  """Lignes : « - » puce (2 espaces = niveau), « 1. » liste numérotée, sinon paragraphe."""
  num=None;out=[]
  for ln in([lines]if isinstance(lines,str)else lines):
   ind=(len(ln)-len(ln.lstrip(" ")))//2;s=ln.strip()
   m=re.match(r"^(?:[-*•]|\d+[.)])\s+(.*)",s)
   if m and not re.match(r"^\d",s):p=self._para(box,m[1],"List Paragraph",ind,self.bullet,size=size,first=first)
   elif m:
    if num is None:num=self._num(self.aid+1,True)
    p=self._para(box,m[1],"List Paragraph",ind,num,size=size,first=first)
   else:num=None;p=self._para(box,s,None if first is not None else None,size=size,first=first)
   first=None;out.append(p)
  return out
 # --- blocs ---
 def block(self,b,i):
  t=b.get("type");f=getattr(self,"b_"+str(t),None)
  if f is None:self.warn.append(f"bloc {i} : type « {t} » inconnu (voir get_word_catalog)");return
  try:f(b)
  except Exception as e:self.warn.append(f"bloc {i} ({t}) : {e}")
 def b_titre(self,b):
  n=min(max(int(b.get("niveau",1)),1),4);p=self.doc.add_paragraph(b.get("texte",""),style=f"Heading {n}")
  if self.hnum:_numpr(p,self.hnum,n-1)  # 1, 1.1, 1.1.1… comme dans le modèle
 def b_paragraphe(self,b):
  st={"mise_en_avant":"Quote","legende":"Caption"}.get(b.get("style",""));self._para(self.doc,b.get("texte",""),st)
 def _items(self,b,mark):
  """Éléments d'une liste : l'indentation (2 espaces = 1 niveau) est conservée, une éventuelle puce saisie est remplacée."""
  out=[]
  for x in b.get("elements",[]):
   x=str(x);ind=len(x)-len(x.lstrip(" "));out.append(" "*ind+mark+re.sub(r"^\s*(?:[-*•]|\d+[.)])?\s*","",x))
  return out
 def b_liste(self,b):self.lines(self.doc,self._items(b,"- "))
 def b_liste_numerotee(self,b):self.lines(self.doc,self._items(b,"1. "))
 def b_citation(self,b):
  self._para(self.doc,b.get("texte",""),"Quote")
  if b.get("auteur"):self._para(self.doc,"— "+b["auteur"],"Quote")
 def b_saut_de_page(self,b):self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
 def b_tableau(self,b):
  rows=b.get("lignes")or[];head=b.get("entetes")or[]
  nc=max([len(head)]+[len(r)for r in rows]);T=self.doc.add_table(rows=0,cols=nc)
  T.style=self.doc.styles[self.cfg["styles_tableau"].get(b.get("style","standard"),self.cfg["styles_tableau"]["standard"])]
  sz=self.cad["taille"]-1
  for ri,r in enumerate(([head]if head else[])+rows):
   cells=T.add_row().cells
   for ci in range(nc):
    c=cells[ci];self._para(c,str(r[ci])if ci<len(r)else"","Tableau",size=sz,first=c.paragraphs[0])
   if ri==0 and head:_repeat_header(T.rows[0])
  if b.get("largeurs"):_width(T,[w*self._usable()/100 for w in b["largeurs"]])
  if b.get("legende"):self._para(self.doc,b["legende"],"Caption")
 def b_encadre(self,b):
  T=self.doc.add_table(rows=1,cols=1);c=T.rows[0].cells[0];_shade(c,self.col["clair"]);_borders(T,self.col["fort"],12,False);_cell_margins(T,100,160);_no_split(T.rows[0])  # encadré jamais coupé
  first=c.paragraphs[0]
  if b.get("titre"):self._para(c,b["titre"],bold=True,color=self.col["fort"],first=first);first=None
  self.lines(c,b.get("contenu")or[],first=first);self.doc.add_paragraph()
 def b_chiffres_cles(self,b):
  ch=b.get("chiffres")or[];T=self.doc.add_table(rows=1,cols=len(ch));_borders(T,"FFFFFF",0,False);_no_split(T.rows[0])
  for c,x in zip(T.rows[0].cells,ch):
   self._para(c,str(x.get("valeur","")),bold=True,color=self.col["fort"],size=self.cad["taille"]+12,first=c.paragraphs[0])
   self._para(c,str(x.get("libelle","")),size=self.cad["taille"]-1)
  self.doc.add_paragraph()
 def b_tableau_risques(self,b):
  L=self.cfg["echelle_risques"];rs=b.get("risques")or[];uo=any(r.get("uo")for r in rs)
  head=["Risque"]+(["UO concernée"]if uo else[])+["Probabilité","Gravité","Criticité","Parade"]
  T=self.doc.add_table(rows=1,cols=len(head));T.style=self.doc.styles[self.cfg["styles_tableau"]["standard"]];sz=self.cad["taille"]-1
  for c,h in zip(T.rows[0].cells,head):self._para(c,h,"Tableau",size=sz,first=c.paragraphs[0])
  _repeat_header(T.rows[0])
  for r in rs:
   p,g=int(r.get("probabilite",2)),int(r.get("gravite",2));k=p*g
   cells=T.add_row().cells;vals=[r.get("risque","")]+([r.get("uo","")]if uo else[])+[L["libelles"][p-1],L["libelles"][g-1],str(k),r.get("parade","")]
   for ci,(c,v) in enumerate(zip(cells,vals)):
    if ci==len(vals)-1 and isinstance(v,list):self.lines(c,v,size=sz,first=c.paragraphs[0])
    else:self._para(c,str(v),"Tableau",size=sz,first=c.paragraphs[0],bold=(ci==len(vals)-2))
   crit=cells[len(vals)-2];_shade(crit,next(x["couleur"]for x in L["criticite"]if k<=x["max"])if self.cad.get("couleurs")!="sobre"else"EDEDED")
  _width(T,[w*self._usable()/100 for w in([28,12,11,11,9,29]if uo else[32,12,12,9,35])])
  self._para(self.doc,b.get("legende")or"Criticité = probabilité × gravité (1 à 16).","Caption")
 def b_fiche_uo(self,b):
  """Fiche d'unité d'œuvre : ligne de titre colorée puis rubriques (titre sur fond clair + contenu)."""
  T=self.doc.add_table(rows=0,cols=1);_borders(T,self.col["fort"],4);_cell_margins(T);sz=self.cad["taille"]-0.5
  c=T.add_row().cells[0];_shade(c,self.col["fort"]);_no_split(T.rows[-1])
  self._para(c,(f"{b['code']} – "if b.get("code")else"")+b.get("titre",""),bold=True,color=self.col.get("texte","FFFFFF"),size=self.cad["taille"]+1,first=c.paragraphs[0])
  for r in b.get("rubriques")or[]:
   h=T.add_row().cells[0];_shade(h,self.col["clair"]);self._para(h,r.get("titre",""),bold=True,size=sz,first=h.paragraphs[0]);_no_split(T.rows[-1])
   cc=T.add_row().cells[0]
   if r.get("tableau"):
    self.lines(cc,r.get("contenu")or[],size=sz,first=cc.paragraphs[0])if r.get("contenu")else None
    rows=r["tableau"];t2=cc.add_table(rows=0,cols=max(map(len,rows)));t2.style=self.doc.styles[self.cfg["styles_tableau"]["standard"]]
    for ri,row in enumerate(rows):
     for x,v in zip(t2.add_row().cells,row):self._para(x,str(v),"Tableau",size=sz-1,first=x.paragraphs[0])
   else:self.lines(cc,r.get("contenu")or["—"],size=sz,first=cc.paragraphs[0])
  self.doc.add_paragraph()
 def b_fiche_profil(self,b):
  T=self.doc.add_table(rows=1,cols=2);_borders(T,self.col["clair"],4);_cell_margins(T,80,140);l,r=T.rows[0].cells;_shade(l,self.col["clair"]);sz=self.cad["taille"]-0.5
  self._para(l,b.get("nom",""),bold=True,color=self.col["fort"],size=self.cad["taille"]+2,first=l.paragraphs[0])
  for x in(b.get("role"),b.get("experience")):
   if x:self._para(l,x,size=sz)
  if b.get("certifications"):self._para(l,"Certifications",bold=True,size=sz);self.lines(l,["- "+x for x in b["certifications"]],size=sz)
  first=r.paragraphs[0]
  for k,lab in(("resume",None),("competences","Compétences clés"),("experiences","Expériences significatives"),("formations","Formations")):
   v=b.get(k)
   if not v:continue
   if lab:self._para(r,lab,bold=True,color=self.col["fort"],size=sz,first=first);first=None
   if isinstance(v,str):self._para(r,v,size=sz,first=first)
   else:self.lines(r,["- "+(x if isinstance(x,str)else f"**{x.get('periode','')} – {x.get('client','')}** : {x.get('mission','')}")for x in v],size=sz,first=first)
   first=None
  _width(T,[self._usable()*.32,self._usable()*.68]);self.doc.add_paragraph()
 def b_presentation_niji(self,b):
  for e in self.niji:self.final.addprevious(deepcopy(e))
 def _usable(self):
  s=self.doc.sections[0];return(s.page_width-s.left_margin-s.right_margin)/360000
 # --- fin : annexes, cadre ---
 def finish(self,opt):
  if opt.get("cgv",False):
   for e in self.tail:self.final.addprevious(e)
  else:self.final.addprevious(self.sect0)
  apply_cadre(self.doc,self.cad,self.cfg)
  return self.doc

def _esc(s):return str(s).replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")

def apply_cadre(doc,cad,cfg):
 """Polices, tailles, interligne, marges, format, couleurs : sur les styles, le thème et les sections."""
 ft,fh=cad.get("police_texte"),cad.get("police_titres")or cad.get("police_texte")
 st=doc.styles;T=cfg["tailles_modele"];base=float(cad["taille"])
 def setf(style,font,size=None,color=None):
  try:s=st[style]
  except KeyError:return
  if font:
   s.font.name=font;rf=s.element.rPr.find(qn("w:rFonts"))
   for a in("w:ascii","w:hAnsi","w:cs","w:eastAsia"):rf.set(qn(a),font)
   for a in("w:asciiTheme","w:hAnsiTheme","w:cstheme","w:eastAsiaTheme"):rf.attrib.pop(qn(a),None)
  if size:s.font.size=Pt(size)
  if color:s.font.color.rgb=RGBColor.from_string(color)
 setf("Normal",ft,base)
 for k,v in T["titres"].items():setf(k,fh,base+v,cad.get("couleur_titres"))
 for k in("List Paragraph","Tableau","Caption","Quote","toc 1","toc 2","toc 3"):setf(k,ft)
 setf("TOC Heading",fh,None,cad.get("couleur_titres"))
 th=doc.part.part_related_by("http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
 if cad.get("couleurs")=="sobre":th._blob=re.sub(rb'(<a:accent4>\s*<a:srgbClr val=")\w{6}',rb"\g<1>595959",th.blob)
 for s in doc.sections:  # pied de page : aligné à gauche (une police plus large ne doit pas l'étirer)
  for p in s.footer.paragraphs:
   if p.alignment is None or int(p.alignment)==3:p.alignment=0
 if cad.get("strict"):  # cadre client : toute police explicite du document suit le cadre (hors puces)
  parts=[doc.part]+[s.header.part for s in doc.sections]+[s.footer.part for s in doc.sections]
  for part in{id(p):p for p in parts}.values():
   for rf in part.element.iter(qn("w:rFonts")):
    if rf.get(qn("w:ascii"))in SYMB:continue
    for a in("w:ascii","w:hAnsi","w:cs"):rf.set(qn(a),ft)
    for a in("w:asciiTheme","w:hAnsiTheme","w:cstheme"):rf.attrib.pop(qn(a),None)
  x=th.blob.decode("utf8");x=re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*',lambda m:m[1]+ft,x);th._blob=x.encode("utf8")
 pf=st["Normal"].paragraph_format
 if cad.get("interligne"):pf.line_spacing=float(cad["interligne"])
 if cad.get("justifie")is not None:
  from docx.enum.text import WD_ALIGN_PARAGRAPH as A;pf.alignment=A.JUSTIFY if cad["justifie"]else A.LEFT
 m=cad.get("marges_cm");W,H=PAGES.get(cad.get("format","A4"),PAGES["A4"])
 for s in doc.sections:
  if cad.get("format"):s.page_width,s.page_height=Cm(W),Cm(H)
  if m is not None:
   mm=m if isinstance(m,dict)else{"haut":m,"bas":m,"gauche":m,"droite":m}
   s.top_margin,s.bottom_margin,s.left_margin,s.right_margin=(Cm(float(mm.get(k,2.5)))for k in("haut","bas","gauche","droite"))
   s.header_distance=min(s.header_distance or Cm(1),s.top_margin);s.footer_distance=min(s.footer_distance or Cm(1),s.bottom_margin)

def count_pages(path,doc,cad):
 """Nombre de pages : exact si LibreOffice (soffice) est présent, sinon estimation (± 15 %)."""
 so=shutil.which("soffice")or shutil.which("libreoffice")
 if so:
  try:
   with tempfile.TemporaryDirectory()as d:
    subprocess.run([so,"--headless","--convert-to","pdf","--outdir",d,str(path)],capture_output=True,timeout=180)
    pdf=next(Path(d).glob("*.pdf"));return len(re.findall(rb"/Type\s*/Page[^s]",pdf.read_bytes())),True
  except Exception:pass
 s=doc.sections[0];w=(s.page_width-s.left_margin-s.right_margin)/12700;h=(s.page_height-s.top_margin-s.bottom_margin)/12700
 z=float(cad["taille"]);li=float(cad.get("interligne")or 1.1);cpl=lambda width:max(10,width/(z*.5));lh=z*1.22*li
 pages,y=1,0.0
 def add(height):
  nonlocal pages,y
  y+=height
  while y>h:pages+=1;y-=h
 body=doc.element.body
 for e in body:
  if e.tag==qn("w:p"):
   t="".join(x.text or""for x in e.iter(qn("w:t")));x=e.xml
   if"<w:sectPr"in x or'w:type="page"'in x:
    if t.strip():add(math.ceil(len(t)/cpl(w))*lh+6)
    pages+=1;y=0;continue
   for d in e.iter(qn("wp:extent")):add(int(d.get("cy"))/12700)
   if t.strip()or"<w:drawing"not in x:add(max(1,math.ceil(len(t)/cpl(w)))*lh+(z*1.6 if"Heading"in x or"Titre"in x else 4))
  elif e.tag==qn("w:tbl"):
   for tr in e.iter(qn("w:tr")):
    tcs=tr.findall(qn("w:tc"))or[None];cw=w/len(tcs)
    add(max(sum(max(1,math.ceil(len("".join(t.text or""for t in p.iter(qn("w:t"))))/cpl(cw-8)))for p in(tc.findall(qn("w:p"))if tc is not None else[]))*lh*.95 for tc in tcs)+6)
  elif e.tag==qn("w:sdt"):add(len(list(e.iter(qn("w:p"))))*lh)
 return pages,False

# ---------- bibliothèque d'UO extraite des réponses passées ----------
_C=r"[A-Z]{2,}(?:-[A-Z0-9]+)+"
UO_TITLE=re.compile(rf"^\s*(UO\s*[\w.-]+|{_C}(?:\s*(?:,|&|et)\s*{_C})*|P\d+(?:\.\d+)*(?:\s*à\s*P?\d+(?:\.\d+)*)?)\s*(?::|\s[–—-]\s)\s*(.+)$")
def extract_uo(paths):
 """Fiches UO = tableaux à une colonne dont la 1re ligne est « CODE : titre » et les rubriques sur fond coloré."""
 lib=[]
 for f in paths:
  d=Document(str(f));src=Path(f).stem
  for t in d.tables:
   if len(t.columns)!=1 or len(t.rows)<3:continue
   head=t.rows[0].cells[0].text.strip();m=UO_TITLE.match(head.replace("\xa0"," "))
   if not m:continue
   rub=[];cur=None
   for r in t.rows[1:]:
    c=r.cells[0];shd=c._tc.find(".//"+qn("w:shd"));fill=shd.get(qn("w:fill"))if shd is not None else None
    ps=[p for p in c.paragraphs if p.text.strip()]
    if fill and fill not in("auto","FFFFFF")and len(ps)<=1:cur={"titre":c.text.strip(),"contenu":[]};rub.append(cur);continue
    if cur is None:cur={"titre":"Description","contenu":[]};rub.append(cur)
    for p in ps:
     lvl=p._p.find(qn("w:pPr")+"/"+qn("w:numPr")+"/"+qn("w:ilvl"))
     ind=int(lvl.get(qn("w:val")))if lvl is not None else None
     cur["contenu"].append(("  "*ind+"- "if ind is not None else"")+p.text.strip().replace("\xa0"," "))
   import unicodedata;i=re.sub(r"[^a-z0-9]+","_",unicodedata.normalize("NFKD",f"{src[:24]}_{m[1]}").encode("ascii","ignore").decode().lower()).strip("_");n=sum(1 for x in lib if x["id"].rsplit("_v",1)[0]==i)
   lib.append({"id":i+(f"_v{n+1}"if n else""),"source":src,"code":m[1].strip(),"titre":m[2].strip(),"rubriques":[x for x in rub if x["titre"]]})
 return lib
